import json
import csv
from decimal import Decimal
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.urls import reverse

from finance.models import FinanceStock
from energy.models import EnergyStock
from basic_idn.models import BasicIndustryStock
from cyclical.models import CyclicalStock
from non_cyclical.models import NonCyclicalStock
from industrial.models import IndustrialStock
from health.models import HealthStock
from infrastructure.models import InfrastructureStock
from property.models import PropertyStock
from technology.models import TechnologyStock
from transport.models import TransportStock
from forex.models import ForexCandle
from watchlist.models import WatchlistItem

from .analytics import calculate_technical_indicators
from finance.management.commands.sync_yahoo_finance import sync_stock_ticker, sync_forex_pair, FOREX_PAIRS

SECTOR_CONFIG = [
    {
        'name': 'Finance',
        'code': 'finance',
        'model': FinanceStock,
        'badge': 'primary',
        'icon': 'bi-bank',
        'desc': 'Perbankan, asuransi, pembiayaan, dan sekuritas.',
        'examples': ['BBCA', 'BBRI', 'BMRI', 'BBNI'],
        'api_url': '/api/finance/',
    },
    {
        'name': 'Energy',
        'code': 'energy',
        'model': EnergyStock,
        'badge': 'warning',
        'icon': 'bi-lightning-charge-fill',
        'desc': 'Minyak, gas bumi, batu bara, dan energi terbarukan.',
        'examples': ['ADRO', 'MEDC', 'PGAS', 'PTBA'],
        'api_url': '/api/energy/',
    },
    {
        'name': 'Basic Materials',
        'code': 'basic-idn',
        'model': BasicIndustryStock,
        'badge': 'secondary',
        'icon': 'bi-gem',
        'desc': 'Pertambangan mineral, logam, semen, dan kimia dasar.',
        'examples': ['ANTM', 'MDKA', 'INCO', 'BRPT'],
        'api_url': '/api/basic-idn/',
    },
    {
        'name': 'Consumer Cyclicals',
        'code': 'cyclical',
        'model': CyclicalStock,
        'badge': 'info',
        'icon': 'bi-cart-fill',
        'desc': 'Ritel, otomotif, pakaian jadi, dan rekreasi.',
        'examples': ['ACES', 'MAPI', 'ERAA'],
        'api_url': '/api/cyclical/',
    },
    {
        'name': 'Consumer Non-Cyclicals',
        'code': 'non-cyclical',
        'model': NonCyclicalStock,
        'badge': 'success',
        'icon': 'bi-cup-straw',
        'desc': 'Makanan, minuman, tembakau, dan kebutuhan pokok.',
        'examples': ['ICBP', 'INDF', 'UNVR', 'MYOR'],
        'api_url': '/api/non-cyclical/',
    },
    {
        'name': 'Industrials',
        'code': 'industrial',
        'model': IndustrialStock,
        'badge': 'dark',
        'icon': 'bi-gear-wide-connected',
        'desc': 'Mesin berat, konstruksi, dan jasa komersial.',
        'examples': ['ASII', 'UNTR', 'HEXA'],
        'api_url': '/api/industrial/',
    },
    {
        'name': 'Healthcare',
        'code': 'health',
        'model': HealthStock,
        'badge': 'danger',
        'icon': 'bi-heart-pulse-fill',
        'desc': 'Farmasi, rumah sakit, dan perlengkapan medis.',
        'examples': ['KLBF', 'MIKA', 'SIDO'],
        'api_url': '/api/health/',
    },
    {
        'name': 'Infrastructure',
        'code': 'infrastructure',
        'model': InfrastructureStock,
        'badge': 'primary',
        'icon': 'bi-broadcast-pin',
        'desc': 'Telekomunikasi, menara pemancar, jalan tol, dan utilitas.',
        'examples': ['TLKM', 'ISAT', 'JSMR', 'TOWR'],
        'api_url': '/api/infrastructure/',
    },
    {
        'name': 'Property & Real Estate',
        'code': 'property',
        'model': PropertyStock,
        'badge': 'warning',
        'icon': 'bi-building',
        'desc': 'Pengembang perumahan, real estat, dan kawasan industri.',
        'examples': ['BSDE', 'CTRA', 'PWON', 'SMRA'],
        'api_url': '/api/property/',
    },
    {
        'name': 'Technology',
        'code': 'technology',
        'model': TechnologyStock,
        'badge': 'info',
        'icon': 'bi-cpu-fill',
        'desc': 'Platform digital, e-commerce, perangkat lunak, dan IT.',
        'examples': ['GOTO', 'BUKA', 'EMTK'],
        'api_url': '/api/technology/',
    },
    {
        'name': 'Transportation & Logistics',
        'code': 'transport',
        'model': TransportStock,
        'badge': 'secondary',
        'icon': 'bi-truck',
        'desc': 'Maskapai penerbangan, armada taksi, logistik, dan kurir.',
        'examples': ['BIRD', 'GIAA', 'ASSA'],
        'api_url': '/api/transport/',
    },
    {
        'name': 'Forex',
        'code': 'forex',
        'model': ForexCandle,
        'badge': 'success',
        'icon': 'bi-currency-exchange',
        'desc': 'Pasangan nilai tukar mata uang asing dan Rupiah.',
        'examples': ['USDIDR', 'EURUSD', 'USDJPY'],
        'api_url': '/api/forex/',
    },
]

def get_stock_by_ticker(ticker_clean):
    """Mencari objek model dan histori data saham berdasarkan ticker di semua sektor."""
    for s in SECTOR_CONFIG:
        if s['code'] == 'forex':
            continue
        qs = s['model'].objects.filter(ticker=ticker_clean).order_by('-date')[:30]
        if qs.exists():
            return s, list(qs)
    # Default fallback ke finance
    return SECTOR_CONFIG[0], list(FinanceStock.objects.filter(ticker=ticker_clean).order_by('-date')[:30])

def home_view(request):
    """Halaman Beranda Tekno-Market Hub dengan statistik dan katalog sektor."""
    total_records = 0
    all_tickers = set()

    sectors_with_stats = []
    for s in SECTOR_CONFIG:
        model = s['model']
        count = model.objects.count()
        total_records += count
        if s['code'] == 'forex':
            tickers = list(model.objects.values_list('pair', flat=True).distinct())
        else:
            tickers = list(model.objects.values_list('ticker', flat=True).distinct())
        all_tickers.update(tickers)
        
        sectors_with_stats.append({
            **s,
            'count': count,
            'active_tickers': tickers,
        })

    context = {
        'sectors': sectors_with_stats,
        'total_sectors': len(SECTOR_CONFIG),
        'total_tickers': len(all_tickers),
        'total_records': total_records,
    }
    return render(request, 'home.html', context)

def stock_dashboard_view(request, ticker):
    """Dashboard visual universal lengkap dengan indikator teknikal (SMA, RSI, BB)."""
    ticker_clean = ticker.upper().strip()
    found_sector, stocks = get_stock_by_ticker(ticker_clean)

    # Cek apakah sudah ada di watchlist
    in_watchlist = False
    session_key = request.session.session_key
    if request.user.is_authenticated:
        in_watchlist = WatchlistItem.objects.filter(user=request.user, ticker=ticker_clean).exists()
    elif session_key:
        in_watchlist = WatchlistItem.objects.filter(session_key=session_key, ticker=ticker_clean).exists()

    latest_close = None
    change = Decimal('0.00')
    change_pct = Decimal('0.00')
    high_30 = None
    low_30 = None
    total_volume = 0
    table_rows = []

    if stocks:
        latest = stocks[0]
        latest_close = latest.close_price
        
        if len(stocks) > 1:
            prev_close = stocks[1].close_price
            change = latest.close_price - prev_close
            change_pct = (change / prev_close * 100) if prev_close else Decimal('0.00')
        else:
            change = latest.close_price - latest.open_price
            change_pct = (change / latest.open_price * 100) if latest.open_price else Decimal('0.00')

        high_30 = max(s.high_price for s in stocks)
        low_30 = min(s.low_price for s in stocks)
        total_volume = sum(s.volume for s in stocks)

        for i, s in enumerate(stocks):
            if i + 1 < len(stocks):
                prev = stocks[i + 1].close_price
                diff = s.close_price - prev
                pct = (diff / prev * 100) if prev else Decimal('0.00')
            else:
                diff = s.close_price - s.open_price
                pct = (diff / s.open_price * 100) if s.open_price else Decimal('0.00')

            table_rows.append({
                'date': s.date,
                'open_price': s.open_price,
                'high_price': s.high_price,
                'low_price': s.low_price,
                'close_price': s.close_price,
                'volume': s.volume,
                'change': diff,
                'change_pct': pct,
                'is_up': diff > 0,
                'is_down': diff < 0,
            })

    # Hitung Indikator Teknikal (SMA, Bollinger Bands, RSI)
    tech = calculate_technical_indicators(stocks)

    # Data Chart.js (kronologis urut tanggal naik)
    chart_stocks = sorted(stocks, key=lambda s: s.date)
    chart_dates = [s.date.strftime('%d %b %Y') for s in chart_stocks]
    chart_closes = [float(s.close_price) for s in chart_stocks]
    chart_volumes = [s.volume for s in chart_stocks]

    peer_tickers = list(found_sector['model'].objects.values_list('ticker', flat=True).distinct())

    context = {
        'ticker': ticker_clean,
        'sector': found_sector,
        'stocks': stocks,
        'table_rows': table_rows,
        'latest_close': latest_close,
        'change': change,
        'change_pct': change_pct,
        'is_up': change > 0,
        'is_down': change < 0,
        'high_30': high_30,
        'low_30': low_30,
        'total_volume': total_volume,
        'peer_tickers': peer_tickers,
        'in_watchlist': in_watchlist,
        # Indikator Teknikal & Rekomendasi Trading
        'tech': tech,
        'recommendation': tech.get('recommendation', {}),
        'analytics': tech.get('recommendation', {}),
        'last_updated': stocks[0].date.strftime('%d %b %Y') if stocks else "Sistem Otomatis",
        # Data List Python Asli untuk filter |json_script di template
        'chart_dates': chart_dates,
        'chart_closes': chart_closes,
        'chart_volumes': chart_volumes,
        'chart_sma5': tech.get('sma_5', []),
        'chart_sma20': tech.get('sma_20', []),
        'chart_ema9': tech.get('ema_9', []),
        'chart_bbupper': tech.get('bb_upper', []),
        'chart_bblower': tech.get('bb_lower', []),
        'chart_dates_json': chart_dates,
        'chart_closes_json': chart_closes,
        'chart_volumes_json': chart_volumes,
        'chart_sma5_json': tech.get('sma_5', []),
        'chart_sma20_json': tech.get('sma_20', []),
        'chart_ema9_json': tech.get('ema_9', []),
        'chart_bbupper_json': tech.get('bb_upper', []),
        'chart_bblower_json': tech.get('bb_lower', []),
        'api_url': f"{found_sector['api_url']}{ticker_clean}/",
    }
    return render(request, 'dashboard.html', context)

def sync_ticker_view(request, ticker):
    """Sinkronisasi data langsung dari Yahoo Finance via tombol di dashboard atau endpoint."""
    ticker_clean = ticker.upper().strip()
    is_forex = ticker_clean in FOREX_PAIRS or '=X' in ticker_clean
    clean_target = ticker_clean.replace('=X', '') if is_forex else ticker_clean

    try:
        if is_forex:
            count = sync_forex_pair(clean_target, period='1mo')
            asset_label = f"Forex {clean_target}"
        else:
            count = sync_stock_ticker(clean_target, period='1mo')
            asset_label = f"Saham {clean_target}"

        if count > 0:
            messages.success(request, f"Sukses! {count} data harga terbaru {asset_label} berhasil disinkronkan dari Yahoo Finance.")
        else:
            messages.warning(request, f"Tidak ada data baru yang ditemukan di Yahoo Finance untuk {asset_label}.")
    except Exception as e:
        messages.error(request, f"Gagal mengambil data dari Yahoo Finance: {e}")

    referer = request.META.get('HTTP_REFERER')
    if referer:
        return redirect(referer)
    if is_forex:
        return redirect(f"/api/forex/{clean_target}/")
    return redirect('stock_dashboard', ticker=clean_target)

def compare_view(request):
    """Fitur perbandingan kinerja multi-emiten saham."""
    tickers_param = request.GET.get('tickers', 'BBCA,BBRI,BMRI')
    ticker_list = [t.strip().upper() for t in tickers_param.split(',') if t.strip()][:4]
    if not ticker_list:
        ticker_list = ['BBCA', 'BBRI']

    comparison_data = []
    all_dates = set()
    ticker_series = {}

    for t in ticker_list:
        sec, stocks = get_stock_by_ticker(t)
        if stocks:
            sorted_s = sorted(stocks, key=lambda s: s.date)
            base_price = float(sorted_s[0].close_price) if sorted_s else 1.0
            
            series = {}
            for s in sorted_s:
                d_str = s.date.strftime('%d %b %Y')
                all_dates.add(s.date)
                pct_change = ((float(s.close_price) - base_price) / base_price) * 100
                series[s.date] = round(pct_change, 2)

            latest = stocks[0]
            first = sorted_s[0]
            period_return = ((float(latest.close_price) - float(first.close_price)) / float(first.close_price)) * 100

            comparison_data.append({
                'ticker': t,
                'sector': sec['name'],
                'latest_price': latest.close_price,
                'period_return': round(period_return, 2),
                'high_30': max(s.high_price for s in stocks),
                'low_30': min(s.low_price for s in stocks),
                'volume': sum(s.volume for s in stocks),
            })
            ticker_series[t] = series

    sorted_dates = sorted(list(all_dates))
    date_labels = [d.strftime('%d %b %Y') for d in sorted_dates]

    # Buat dataset Chart.js
    COLORS = ['#0d6efd', '#198754', '#dc3545', '#ffc107', '#6f42c1']
    datasets = []
    for idx, t in enumerate(ticker_list):
        series = ticker_series.get(t, {})
        data_points = []
        last_val = 0
        for d in sorted_dates:
            if d in series:
                last_val = series[d]
            data_points.append(last_val)
        
        datasets.append({
            'label': t,
            'data': data_points,
            'borderColor': COLORS[idx % len(COLORS)],
            'backgroundColor': 'transparent',
            'borderWidth': 2.5,
            'tension': 0.2,
        })

    # Dapatkan semua pilihan ticker untuk selectbox
    available_tickers = []
    for s in SECTOR_CONFIG:
        if s['code'] != 'forex':
            available_tickers.extend(list(s['model'].objects.values_list('ticker', flat=True).distinct()))

    context = {
        'selected_tickers': ticker_list,
        'selected_tickers_str': ','.join(ticker_list),
        'comparison_data': comparison_data,
        'date_labels_json': json.dumps(date_labels),
        'datasets_json': json.dumps(datasets),
        'available_tickers': sorted(list(set(available_tickers))),
    }
    return render(request, 'compare.html', context)

def export_csv_view(request, ticker):
    """Mengekspor histori harga saham ke file CSV."""
    ticker_clean = ticker.upper().strip()
    sec, stocks = get_stock_by_ticker(ticker_clean)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{ticker_clean}_histori_data.csv"'

    writer = csv.writer(response)
    writer.writerow(['Ticker', 'Date', 'Open', 'High', 'Low', 'Close', 'Volume'])

    for s in stocks:
        writer.writerow([s.ticker, s.date.strftime('%Y-%m-%d'), s.open_price, s.high_price, s.low_price, s.close_price, s.volume])

    return response

def export_excel_view(request, ticker):
    """Mengekspor histori harga saham ke file Excel (.xlsx)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    ticker_clean = ticker.upper().strip()
    sec, stocks = get_stock_by_ticker(ticker_clean)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{ticker_clean} History"

    # Header Styling
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = ['Ticker', 'Tanggal', 'Open (Rp)', 'High (Rp)', 'Low (Rp)', 'Close (Rp)', 'Volume']
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for s in stocks:
        ws.append([
            s.ticker,
            s.date.strftime('%Y-%m-%d'),
            float(s.open_price),
            float(s.high_price),
            float(s.low_price),
            float(s.close_price),
            int(s.volume),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{ticker_clean}_histori_data.xlsx"'
    wb.save(response)
    return response

def watchlist_view(request):
    """Halaman Watchlist Portofolio Pribadi & Pantauan Sinyal Harga."""
    session_key = request.session.session_key
    if not session_key:
        request.session.create()
        session_key = request.session.session_key

    if request.user.is_authenticated:
        items = WatchlistItem.objects.filter(user=request.user)
    else:
        items = WatchlistItem.objects.filter(session_key=session_key)

    enriched_items = []
    for item in items:
        sec, stocks = get_stock_by_ticker(item.ticker)
        latest_price = None
        change_pct = Decimal('0.00')
        tech_signal = 'Netral'
        tech_badge = 'secondary'

        if stocks:
            latest = stocks[0]
            latest_price = latest.close_price
            if len(stocks) > 1:
                prev = stocks[1].close_price
                change_pct = ((latest.close_price - prev) / prev * 100) if prev else Decimal('0.00')
            
            tech = calculate_technical_indicators(stocks)
            tech_signal = tech['rsi_signal']
            rec = tech.get('recommendation', {})
        else:
            rec = {}

        enriched_items.append({
            'item': item,
            'sector': sec,
            'latest_price': latest_price,
            'change_pct': change_pct,
            'is_up': change_pct > 0,
            'is_down': change_pct < 0,
            'tech_signal': tech_signal,
            'tech_badge': tech_badge,
            'rec': rec,
        })

    context = {
        'watchlist_items': enriched_items,
    }
    return render(request, 'watchlist.html', context)

def toggle_watchlist_view(request, ticker):
    """Menambahkan atau menghapus saham dari Watchlist pengguna."""
    ticker_clean = ticker.upper().strip()
    sec, _ = get_stock_by_ticker(ticker_clean)

    session_key = request.session.session_key
    if not session_key:
        request.session.create()
        session_key = request.session.session_key

    if request.user.is_authenticated:
        obj = WatchlistItem.objects.filter(user=request.user, ticker=ticker_clean).first()
        if obj:
            obj.delete()
            messages.info(request, f"{ticker_clean} dihapus dari Watchlist Anda.")
        else:
            WatchlistItem.objects.create(user=request.user, ticker=ticker_clean, sector=sec['name'])
            messages.success(request, f"{ticker_clean} berhasil ditambahkan ke Watchlist!")
    else:
        obj = WatchlistItem.objects.filter(session_key=session_key, ticker=ticker_clean).first()
        if obj:
            obj.delete()
            messages.info(request, f"{ticker_clean} dihapus dari Watchlist Anda.")
        else:
            WatchlistItem.objects.create(session_key=session_key, ticker=ticker_clean, sector=sec['name'])
            messages.success(request, f"{ticker_clean} berhasil ditambahkan ke Watchlist!")

    return redirect('stock_dashboard', ticker=ticker_clean)

def openapi_spec_view(request):
    """Mengembalikan spesifikasi OpenAPI 3.0 JSON untuk semua 12 sektor pasar."""
    paths = {}
    for s in SECTOR_CONFIG:
        if s['code'] == 'forex':
            path_url = f"/api/forex/{{pair}}/"
            paths[path_url] = {
                'get': {
                    'summary': f"Data Nilai Tukar Forex ({s['name']})",
                    'description': f"Mengambil data harga bid, ask, dan volume nilai tukar mata uang.",
                    'parameters': [{
                        'name': 'pair',
                        'in': 'path',
                        'required': True,
                        'schema': {'type': 'string', 'example': 'USDIDR'},
                        'description': 'Kode pasangan mata uang, misal USDIDR, EURUSD, USDJPY'
                    }],
                    'responses': {
                        '200': {'description': 'Sukses mengembalikan histori kurs valas'}
                    }
                }
            }
        else:
            path_url = f"{s['api_url']}{{ticker}}/"
            paths[path_url] = {
                'get': {
                    'summary': f"Data Saham Sektor {s['name']}",
                    'description': f"Mengambil 30 data histori OHLC saham sektor {s['name']}.",
                    'parameters': [{
                        'name': 'ticker',
                        'in': 'path',
                        'required': True,
                        'schema': {'type': 'string', 'example': s['examples'][0]},
                        'description': f"Kode ticker saham, misal: {', '.join(s['examples'])}"
                    }],
                    'responses': {
                        '200': {'description': 'Sukses mengembalikan histori data saham'}
                    }
                }
            }

    spec = {
        'openapi': '3.0.3',
        'info': {
            'title': 'Tekno-Market Hub REST API',
            'version': '2.0.0',
            'description': 'Dokumentasi interaktif REST API untuk seluruh 12 sektor pasar saham IHSG dan Valuta Asing.',
        },
        'servers': [{'url': '/'}],
        'paths': paths
    }
    return JsonResponse(spec)

def swagger_docs_view(request):
    """Tampilan UI interaktif Swagger Documentation."""
    return render(request, 'swagger_docs.html')
